import streamlit as st
import pandas as pd
from docx import Document
import sqlite3

# 데이터베이스 연결 함수
def get_connection():
    conn = sqlite3.connect('blr_app/troubleshooting.db', check_same_thread=False)
    return conn

# 데이터베이스 초기화 및 테이블 생성
def create_table(table_name, headers):
    conn = get_connection()
    c = conn.cursor()
    
    # 테이블 이름과 컬럼 동적으로 생성
    columns = ', '.join([f'"{header}" TEXT' for header in headers])  # 헤더를 컬럼으로 사용
    query = f'''
        CREATE TABLE IF NOT EXISTS "{table_name}" (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            {columns}
        )
    '''
    c.execute(query)
    conn.commit()
    conn.close()

# 데이터베이스의 모든 데이터 삭제
def clear_table(table_name):
    conn = get_connection()
    c = conn.cursor()
    c.execute(f'DELETE FROM "{table_name}"')  # 테이블의 모든 데이터 삭제
    conn.commit()
    conn.close()

# 데이터베이스에 데이터 삽입 (중복 체크 포함)
def insert_data(table_name, headers, row):
    conn = get_connection()
    c = conn.cursor()

    placeholders = ', '.join(['?' for _ in headers])  # ? 개수를 headers 길이에 맞게 설정
    columns = ', '.join([f'"{header}"' for header in headers])  # 헤더를 컬럼으로 사용

    # 중복된 내용이 있는지 확인
    check_query = f'SELECT COUNT(*) FROM "{table_name}" WHERE '
    check_conditions = ' AND '.join([f'"{header}" = ?' for header in headers])
    check_query += check_conditions

    if c.execute(check_query, row).fetchone()[0] == 0:  # 중복된 내용이 없으면
        query = f'''
            INSERT INTO "{table_name}" ({columns})
            VALUES ({placeholders})
        '''
        c.execute(query, row)

    conn.commit()
    conn.close()

# 데이터베이스에서 데이터 조회
def view_data(table_name, headers):
    conn = get_connection()
    c = conn.cursor()
    columns = ', '.join([f'"{header}"' for header in headers])  # 헤더를 컬럼으로 사용
    query = f'SELECT {columns} FROM "{table_name}"'
    c.execute(query)
    data = c.fetchall()
    conn.close()
    return data

# Word 파일에서 표 데이터를 읽어오는 함수
def read_word_table(file):
    document = Document(file)
    tables_data = []
    
    # 모든 표를 순회하여 데이터 추출
    for table in document.tables:
        if len(table.rows) > 0:
            # 첫 번째 행을 표 제목으로, 두 번째 행을 헤더로 사용
            table_name = table.rows[0].cells[0].text.strip()  # 첫 번째 셀을 표 이름으로 설정
            header = [cell.text.strip() for cell in table.rows[1].cells]  # 두 번째 행을 헤더로 사용

            # 중복된 헤더가 있을 경우 번호를 붙여서 고유하게 만듦
            unique_header = []
            header_count = {}
            for col in header:
                if col in header_count:
                    header_count[col] += 1
                    unique_header.append(f"{col}_{header_count[col]}")  # 중복된 경우 번호를 추가
                else:
                    header_count[col] = 1
                    unique_header.append(col)

            table_data = []
            for row in table.rows[2:]:  # 나머지 행을 데이터로 사용
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)
            if table_data:  # 데이터가 있는 경우에만 추가
                tables_data.append((table_name, unique_header, table_data))
    
    return tables_data

# Troubleshooting 페이지 함수
def app():
    st.title("Troubleshooting")

    # 관리자 모드 체크박스
    is_admin = st.sidebar.checkbox("관리자 모드 활성화")

    # 데이터 저장 여부를 추적하기 위한 세션 상태
    if 'data_saved' not in st.session_state:
        st.session_state['data_saved'] = False

    # Word 파일 업로드 (관리자 모드에서만 가능)
    if is_admin:
        uploaded_file = st.file_uploader("Word 파일을 업로드하세요", type=["docx"])

        # 파일 업로드된 경우 처리
        if uploaded_file is not None:
            # 데이터베이스의 기존 데이터 삭제
            conn = get_connection()
            c = conn.cursor()
            c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name != 'sqlite_sequence';")
            tables_in_db = c.fetchall()

            # 기존 테이블의 데이터를 삭제
            for table_name_tuple in tables_in_db:
                table_name = table_name_tuple[0]
                clear_table(table_name)

            # Word 파일에서 표 읽기
            tables = read_word_table(uploaded_file)

            # 각 표의 이름과 내용을 저장하고 표시
            for table_name, header, table in tables:
                # 테이블 생성 (헤더를 기반으로)
                create_table(table_name, header)

                # 데이터 삽입
                for row in table:
                    insert_data(table_name, header, row)

            st.session_state['data_saved'] = True
            st.success("파일이 업로드되고 데이터베이스에 저장되었습니다.")

    # 데이터베이스에 저장된 테이블 목록을 불러오기
    conn = get_connection()
    c = conn.cursor()
    c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name != 'sqlite_sequence';")
    tables_in_db = c.fetchall()
    conn.close()

    # 저장된 테이블들 중에서 선택할 수 있도록 표시
    if tables_in_db:
        table_names = [table_name_tuple[0] for table_name_tuple in tables_in_db]

        # 관리자 모드에서만 사이드바에 저장된 파일 목록과 삭제 기능 추가
        if is_admin:
            st.sidebar.title("저장된 파일 목록")
            for table_name in table_names:
                if st.sidebar.button(f"{table_name} 삭제"):
                    conn = get_connection()
                    c = conn.cursor()
                    c.execute(f'DROP TABLE "{table_name}"')
                    conn.commit()
                    conn.close()
                    st.sidebar.success(f"{table_name} 테이블이 삭제되었습니다.")

        selected_table = st.selectbox("Troubleshooting을 선택하시오: ", table_names)

        # 선택한 테이블의 저장된 데이터를 조회하고 화면에 표시
        if selected_table:
            c = get_connection().cursor()
            c.execute(f'PRAGMA table_info("{selected_table}")')
            headers_info = c.fetchall()
            headers = [header_info[1] for header_info in headers_info[1:]]  # id 필드를 제외한 헤더 목록

            stored_data = view_data(selected_table, headers)
            if stored_data:
                st.subheader(f"{selected_table}")
                df = pd.DataFrame(stored_data, columns=headers)  # ID 필드를 제외하고 데이터프레임 생성

                # 데이터 편집 가능하게 표시 (키보드로 값 수정 가능, 관리자 모드에서만 가능)
                if is_admin:
                    edited_df = st.data_editor(df, num_rows="dynamic", key=f'editable_table_{selected_table}')

                    # "변경 사항 저장" 버튼을 통해 수정된 값들을 데이터베이스에 업데이트
                    if st.button(f"변경 사항 저장 ({selected_table})"):
                        conn = get_connection()
                        c = conn.cursor()

                        # 업데이트 쿼리 실행
                        for index, row in edited_df.iterrows():
                            placeholders = ', '.join([f'"{header}" = ?' for header in headers])
                            query = f'UPDATE "{selected_table}" SET {placeholders} WHERE id = ?'
                            values = [row[header] for header in headers] + [row['id']]
                            c.execute(query, values)

                        conn.commit()
                        conn.close()
                        st.success(f"모든 변경 사항이 {selected_table}에 저장되었습니다.")
                else:
                    st.write(df)
            else:
                st.write("선택한 테이블에 저장된 데이터가 없습니다.")
    else:
        st.write("저장된 데이터가 없습니다. Word 파일을 업로드하여 데이터를 추가하세요.")

# Streamlit 앱 실행
if __name__ == "__main__":
    app()
